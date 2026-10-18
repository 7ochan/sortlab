"""Build the editable report DOCX from the reviewed HTML source."""
from pathlib import Path
from lxml import html
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "SortLab-Final-Report.html"
OUTPUT = ROOT / "SortLab-Final-Report.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(32, 39, 51)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 16, 8),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)

def add_inline(paragraph, node, bold=False, italic=False):
    if node.text:
        run = paragraph.add_run(node.text)
        run.bold, run.italic = bold, italic
    for child in node:
        tag = child.tag.lower() if isinstance(child.tag, str) else ""
        b = bold or tag in ("b", "strong")
        i = italic or tag in ("i", "em")
        if tag in ("br",):
            paragraph.add_run().add_break()
        elif tag in ("img",):
            src = child.get("src", "")
            img_path = (ROOT / src).with_suffix(".png")
            if not img_path.exists():
                img_path = ROOT / src
            if img_path.exists():
                paragraph.add_run().add_picture(str(img_path), width=Inches(6.25))
        else:
            add_inline(paragraph, child, b, i)
        if child.tail:
            run = paragraph.add_run(child.tail)
            run.bold, run.italic = bold, italic

def add_paragraph(node, style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    add_inline(p, node)
    return p

def add_table(node):
    source_rows = node.xpath("./tr | ./thead/tr | ./tbody/tr")
    if not source_rows:
        return
    cols = max(len(r.xpath("./th | ./td")) for r in source_rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for ridx, source_row in enumerate(source_rows):
        cells = table.add_row().cells
        cells[0].width = Inches(6.5 / cols)
        if ridx == 0:
            set_repeat_table_header(table.rows[-1])
        for cidx, src_cell in enumerate(source_row.xpath("./th | ./td")):
            cell = cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if src_cell.tag.lower() == "th" or ridx == 0:
                set_cell_shading(cell, "F2F4F7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, src_cell, bold=src_cell.tag.lower() == "th")
            for run in p.runs:
                run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def walk(container, in_cover=False, in_refs=False):
    for node in container:
        tag = node.tag.lower() if isinstance(node.tag, str) else ""
        classes = node.get("class", "").split()
        if tag in ("h1", "h2", "h3"):
            style = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[tag]
            add_paragraph(node, style)
        elif tag == "p":
            p = add_paragraph(node)
            if in_refs:
                p.paragraph_format.space_after = Pt(8)
                for run in p.runs:
                    run.font.size = Pt(9)
        elif tag in ("ul", "ol"):
            for li in node.xpath("./li"):
                style = "List Bullet" if tag == "ul" else "List Number"
                p = doc.add_paragraph(style=style)
                p.paragraph_format.left_indent = Inches(.5)
                p.paragraph_format.first_line_indent = Inches(-.25)
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, li)
        elif tag == "table":
            add_table(node)
        elif tag == "figure":
            cap = " ".join(node.xpath(".//figcaption//text()"))
            if cap.startswith("Figure 2."):
                continue  # same schematic repeated; retain the first workcell diagram
            img = node.find("img")
            if img is not None:
                src = img.get("src", "")
                path = (ROOT / src).with_suffix(".png")
                if not path.exists():
                    path = ROOT / src
                if path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(2)
                    p.add_run().add_picture(str(path), width=Inches(6.25))
            if cap:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                for run in p.add_run(cap),:
                    run.italic = True
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(82, 96, 112)
        elif tag == "div" and "title" in classes:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("".join(node.itertext()).strip())
            run.font.name, run.font.size = "Calibri", Pt(27)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string("1F4D78")
        elif tag == "div" and "subtitle" in classes:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(20)
            run = p.add_run("".join(node.itertext()).strip())
            run.font.size = Pt(17)
            run.font.color.rgb = RGBColor.from_string("435368")
        elif tag == "div" and "kicker" in classes:
            p = doc.add_paragraph()
            run = p.add_run("".join(node.itertext()).strip())
            run.font.size, run.font.bold = Pt(10), True
            run.font.color.rgb = RGBColor.from_string("1F4D78")
        elif tag == "div" and "equation" in classes:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(7)
            add_inline(p, node)
            for run in p.runs:
                run.font.name = "Cambria"
                run.font.size = Pt(11)
        elif tag == "div" and "callout" in classes:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.12)
            p.paragraph_format.right_indent = Inches(.12)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, node)
        elif tag == "section":
            if "refs" in classes:
                doc.add_page_break()
                walk(node, in_refs=True)
            elif "cover" in classes:
                walk(node, in_cover=True)
                doc.add_page_break()
            else:
                walk(node, in_cover=in_cover, in_refs=in_refs)
        elif tag in ("div", "body", "html"):
            if tag == "div" and (node.text or "").strip() and not len(node):
                p = doc.add_paragraph()
                add_inline(p, node)
            else:
                walk(node, in_cover=in_cover, in_refs=in_refs)

tree = html.parse(str(SOURCE))
walk(tree.getroot())

# Header/footer identify the document without adding decorative metrics.
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("SORTLAB  |  FINAL DESIGN REPORT  |  V3.0")
r.font.name = "Calibri"
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(100, 112, 125)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Robotics · BCA UG Semester VII")
r.font.name = "Calibri"
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(100, 112, 125)

doc.core_properties.title = "SortLab — Automated Warehouse Package Sorting Robot — Version 3.0"
doc.core_properties.subject = "Robotics individual design project final report"
doc.core_properties.author = "Student author name to be supplied"
doc.save(OUTPUT)
print(f"Wrote {OUTPUT}")
